"""
Phase 5 Test Suite: Insights Report, Word Document, HTML & Dashboard
Validates PRD Section 5.1-5.7, 6 (FR-9, FR-10) and Phase 5 Exit Criteria.
"""

import os
import unittest
from docx import Document

class TestPhase5Reports(unittest.TestCase):
    def setUp(self):
        self.md_report = os.path.join("04_Reports", "HireSmart_Task1_Hiring_Insights_Report.md")
        self.docx_report = os.path.join("04_Reports", "HireSmart_Task1_Hiring_Insights_Report.docx")
        self.html_report = os.path.join("04_Reports", "HireSmart_Task1_Hiring_Insights_Report.html")
        self.dashboard_file = os.path.join("dashboard", "index.html")

    def test_markdown_report_structure(self):
        """Phase 5 & FR-10: Verify 1-page Markdown report structure and contents."""
        self.assertTrue(os.path.exists(self.md_report), "Markdown report missing.")
        with open(self.md_report, "r", encoding="utf-8") as f:
            content = f.read()

        # Check required sections per PRD Phase 5
        self.assertIn("1. Executive Headline", content, "Section 1 Headline missing.")
        self.assertIn("2. Core Hiring Metrics Scorecard", content, "Section 2 Metrics table missing.")
        self.assertIn("3. Key Findings", content, "Section 3 Findings missing.")
        self.assertIn("4. Ranked Process Bottlenecks", content, "Section 4 Bottlenecks missing.")
        self.assertIn("5. Actionable Handoff to Task 2", content, "Section 5 Next steps missing.")

    def test_docx_report_no_placeholders_left(self):
        """PRD Section 11 & FR-10: Confirm all [from data] placeholders are filled in docx."""
        self.assertTrue(os.path.exists(self.docx_report), "DOCX report missing.")
        doc = Document(self.docx_report)
        
        full_text = []
        for p in doc.paragraphs:
            full_text.append(p.text)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    full_text.append(cell.text)
        
        combined = " ".join(full_text)
        self.assertNotIn("[from data]", combined, "DOCX report still contains unresolved '[from data]' placeholders!")
        self.assertIn("48.0 days", combined, "DOCX report missing 48.0 days TTH calculation.")
        self.assertIn("67.5%", combined, "DOCX report missing 67.5% acceptance rate.")
        self.assertIn("19.6%", combined, "DOCX report missing 19.6% early exit rate.")

    def test_html_report_print_ready(self):
        """FR-10: Verify print-ready HTML executive report exists and contains styled scorecards."""
        self.assertTrue(os.path.exists(self.html_report), "HTML report missing.")
        with open(self.html_report, "r", encoding="utf-8") as f:
            html = f.read()

        self.assertIn("@page", html, "CSS print styling missing.")
        self.assertIn("kpi-card", html, "KPI scorecards missing.")
        self.assertIn("48.0d", html, "Time-to-hire KPI value missing.")
        self.assertIn("charts/offers_vs_acceptance_trend.png", html, "Linked chart image missing.")

    def test_interactive_dashboard_features(self):
        """FR-10: Confirm interactive intelligence dashboard features and simulator."""
        self.assertTrue(os.path.exists(self.dashboard_file), "Dashboard index.html missing.")
        with open(self.dashboard_file, "r", encoding="utf-8") as f:
            dash = f.read()

        self.assertIn("chart.js", dash.lower(), "Chart.js library missing from dashboard.")
        self.assertIn("Automated Screening Impact Simulator", dash, "Task 2 Simulator missing.")
        self.assertIn("updateSim", dash, "Simulator JS interactive logic missing.")

if __name__ == "__main__":
    unittest.main()
